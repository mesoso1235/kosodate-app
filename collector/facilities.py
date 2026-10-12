"""施設データ(世田谷区オープンデータ)の取得と変換。週1回だけ実行される。

- 保育施設・幼稚園・児童館・おでかけひろば・授乳スペース: 保育施設空き情報 xlsx(座標つき)
- 子育てステーション: st.csv(座標なし) → 国土地理院の住所検索で座標化(結果は geocode_cache.json に保存)
- 座標のない行は載せない(地図に出せないため)。
"""
from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path
from urllib.parse import quote, urljoin

import openpyxl

from collector.model import clean_text, item_id, normalize_date

SOURCE_ID = "facilities"
NAME = "世田谷区オープンデータ(施設)"
CATEGORY = "区"
SOURCE = "世田谷区オープンデータ"
LICENSE = "CC BY 4.0（世田谷区）"

PAGE_URL = "https://www.city.setagaya.lg.jp/01044/1577.html"
FALLBACK_XLSX = "https://www.city.setagaya.lg.jp/documents/1577/hoikusisetuakizyoho_1.xlsx"
STATION_CSV = "https://www.city.setagaya.lg.jp/documents/10728/st.csv"
GSI = "https://msearch.gsi.go.jp/address-search/AddressSearch?q="

TYPE_MAP = {
    "私立保育園": "認可保育園",
    "区立保育園": "認可保育園",
    "認証保育所": "認証保育所",
    "小規模保育事業A型": "小規模保育",
    "家庭的保育事業": "小規模保育",
    "事業所内保育事業": "小規模保育",
    "保育ママ": "保育室・保育ママ",
    "保育室": "保育室・保育ママ",
    "企業主導型保育事業": "企業主導型保育",
    "病児・病後児保育": "病児・病後児保育",
    "私立幼稚園": "幼稚園",
    "区立幼稚園": "幼稚園",
    "認定こども園": "認定こども園",
}


# ---------------------------------------------------------------- ページ解析

def find_xlsx_url(html: str) -> str:
    """公式ページから .xlsx へのリンクを探す。ファイル名が変わっても追えるようにする。"""
    links = re.findall(r'href="([^"]+\.xlsx)"', html)
    for href in links:
        if "hoiku" in href.lower() or "documents/1577/" in href:
            return urljoin(PAGE_URL, href)
    return urljoin(PAGE_URL, links[0]) if links else FALLBACK_XLSX


def find_as_of(html: str) -> str | None:
    text = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", html))
    m = re.search(r"最終更新日?\s*(\d{4}年\d{1,2}月\d{1,2}日)", text)
    return normalize_date(m.group(1)) if m else None


# ---------------------------------------------------------------- xlsx

def _s(v) -> str:
    if v is None:
        return ""
    v = clean_text(str(v))
    return "" if v == "-" else v


def _num(v) -> float | None:
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return f if f else None


def _col(header: list, name: str) -> int:
    norm = [re.sub(r"\s", "", str(h)) if h is not None else "" for h in header]
    return norm.index(name)


def _full_address(a) -> str:
    a = _s(a)
    if not a or a.startswith(("東京都", "世田谷区")):
        return a
    return "世田谷区" + a


def _kind_from_name(name: str) -> str:
    if "児童館" in name:
        return "児童館"
    if "ひろば" in name or "広場" in name:
        return "おでかけひろば"
    return "遊び場"


def _name_key(name: str) -> str:
    return re.sub(r"[\s　]", "", name)


_DIGITS = str.maketrans("０１２３４５", "012345")


def parse_xlsx(content: bytes, as_of: str | None = None) -> tuple[list[dict], dict]:
    """(items, stats) を返す。stats は座標なしで落とした件数などの記録。"""
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)

    vac: dict[str, dict] = {}
    if "全地域保育空き状況" in wb.sheetnames:
        rows = list(wb["全地域保育空き状況"].iter_rows(values_only=True))
        hdr = list(rows[0])
        ages = {}
        for i, h in enumerate(hdr):
            m = re.fullmatch(r"([０-５0-5])歳", re.sub(r"\s", "", str(h or "")))
            if m:
                ages[i] = f"{m.group(1).translate(_DIGITS)}歳"
        ni = _col(hdr, "保育施設名称")
        for r in rows[1:]:
            if r[ni]:
                vac[_name_key(str(r[ni]))] = {a: ("" if r[i] is None else str(r[i])) for i, a in ages.items()}

    rows = list(wb["施設一覧"].iter_rows(values_only=True))
    hdr = list(rows[0])
    ix = {k: _col(hdr, k) for k in
          ["施設名", "施設カテゴリ", "保育施設種別", "住所情報", "電話番号", "情報取得元URL", "緯度", "経度"]}

    items, stats = [], {"no_coords": 0, "unknown_type": 0, "vacancy_matched": 0}
    for r in rows[1:]:
        name = _s(r[ix["施設名"]])
        cat = _s(r[ix["施設カテゴリ"]])
        if not name or not cat:
            continue
        kind = _s(r[ix["保育施設種別"]])
        if cat == "幼稚園・保育施設":
            ftype = TYPE_MAP.get(kind)
            if not ftype:
                stats["unknown_type"] += 1
                ftype = kind or "保育施設"
        elif cat == "児童館・おでかけひろば等":
            ftype = _kind_from_name(name)
        elif cat == "おむつ替え・授乳・搾乳スペース":
            ftype = "授乳室"
        else:
            continue
        lat, lng = _num(r[ix["緯度"]]), _num(r[ix["経度"]])
        if lat is None or lng is None:
            stats["no_coords"] += 1
            continue
        address = _full_address(r[ix["住所情報"]])
        it = {
            "id": item_id(f"{name}|{address}"),
            "name": name,
            "type": ftype,
            "address": address,
            "lat": lat,
            "lng": lng,
            "tel": _s(r[ix["電話番号"]]),
            "url": _s(r[ix["情報取得元URL"]]),
            "source": SOURCE,
            "license": LICENSE,
        }
        if "保育" in ftype:
            v = vac.get(_name_key(name))
            if v:
                it["vacancy"] = v
                it["vacancy_as_of"] = as_of
                stats["vacancy_matched"] += 1
        items.append(it)
    return items, stats


# ---------------------------------------------------------------- 子育てステーション

def _tel(text: str) -> str:
    m = re.search(r"\d{2,4}-\d{2,4}-\d{4}", text)
    return m.group(0) if m else ""


def parse_station_csv(content: bytes, geocode) -> tuple[list[dict], int]:
    text = content.decode("cp932", errors="replace")
    groups: list[list[list[str]]] = []
    for row in csv.reader(io.StringIO(text)):
        row = (row + [""] * 4)[:4]
        if row[0].startswith("子育てステーション") and row[0] != "子育てステーション一覧":
            groups.append([row])
        elif groups:
            groups[-1].append(row)
    items, skipped = [], 0
    for g in groups:
        name = clean_text(g[0][0])
        addr = next((clean_text(r[1]) for r in g if clean_text(r[1]).startswith("東京都")), "")
        tel = next((_tel(r[3]) for r in g if _tel(r[3])), "")
        pos = geocode(addr) if addr else None
        if not pos:
            skipped += 1
            continue
        items.append({
            "id": item_id(f"{name}|{addr}"),
            "name": name,
            "type": "子育てステーション",
            "address": addr,
            "lat": pos[0],
            "lng": pos[1],
            "tel": tel,
            "url": "https://www.city.setagaya.lg.jp/02413/10728.html",
            "source": SOURCE,
            "license": LICENSE,
        })
    return items, skipped


class Geocoder:
    """国土地理院の住所検索。結果は cache_path に残し、同じ住所は二度と問い合わせない。"""

    def __init__(self, session, cache_path: Path):
        self.session = session
        self.path = Path(cache_path)
        try:
            self.cache = json.loads(self.path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            self.cache = {}
        self.dirty = False

    @staticmethod
    def variants(address: str) -> list[str]:
        """「東京都世田谷区下馬2丁目37番15号昭和女子大…」→ 建物名を落とした住所の候補。"""
        m = re.match(r"(.*?\d+丁目\d+番\d+号|.*?\d+丁目\d+番|.*?\d+丁目|.*?\d+-\d+-\d+|.*?\d+-\d+|.*?\d+)", address)
        base = m.group(1) if m else address
        hy = re.sub(r"(\d+)丁目(\d+)番(\d+)号", r"\1-\2-\3", base)
        hy = re.sub(r"(\d+)丁目(\d+)番", r"\1-\2", hy)
        out: list[str] = []
        for v in (hy, base, address):
            if v and v not in out:
                out.append(v)
        return out

    def __call__(self, address: str):
        if address in self.cache:
            v = self.cache[address]
            return tuple(v) if v else None
        for q in self.variants(address):
            try:
                data = self.session.get(GSI + quote(q)).json()
            except Exception:
                return None  # 一時的な失敗は保存しない
            if data:
                lng, lat = data[0]["geometry"]["coordinates"][:2]
                self.cache[address] = [lat, lng]
                self.dirty = True
                return (lat, lng)
        self.cache[address] = None
        self.dirty = True
        return None

    def save(self) -> None:
        if self.dirty:
            self.path.write_text(json.dumps(self.cache, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")


def fetch(session, cache_path: Path) -> tuple[list[dict], dict]:
    page = session.get(PAGE_URL).content.decode("utf-8", errors="replace")
    items, stats = parse_xlsx(session.get(find_xlsx_url(page)).content, find_as_of(page))
    geo = Geocoder(session, cache_path)
    try:
        st, skipped = parse_station_csv(session.get(STATION_CSV).content, geo)
        items += st
        stats["station_no_coords"] = skipped
    finally:
        geo.save()
    return items, stats
